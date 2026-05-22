"""简历类文档识别：图片 OCR；PDF/Word 优先抽文本，扫描 PDF 回退分页 OCR。"""

from __future__ import annotations

import logging
import re
import tempfile
from pathlib import Path
from typing import Any

from .core import run_ocr
from .models import OCRLine, OCRSettings

logger = logging.getLogger(__name__)

ALLOWED_IMAGE_SUFFIX = {".png", ".jpg", ".jpeg", ".bmp", ".webp", ".tif", ".tiff"}
ALLOWED_DOCUMENT_SUFFIX = {".pdf", ".docx"}
ALLOWED_FILE_SUFFIX = ALLOWED_IMAGE_SUFFIX | ALLOWED_DOCUMENT_SUFFIX

# 抽出的纯文本少于此阈值时，视为扫描件 PDF，改走分页 OCR
MIN_PDF_TEXT_CHARS = 50
MAX_PDF_PAGES = 15
PDF_RENDER_SCALE = 2.0


def is_allowed_suffix(suffix: str) -> bool:
    s = (suffix or "").lower()
    if not s:
        return True
    return s in ALLOWED_FILE_SUFFIX


def _meaningful_char_count(text: str) -> int:
    return len(re.sub(r"\s+", "", text or ""))


def _lines_from_plain_text(text: str, *, score: float = 1.0) -> list[OCRLine]:
    lines: list[OCRLine] = []
    for raw in (text or "").splitlines():
        t = raw.strip()
        if t:
            lines.append(OCRLine(text=t, score=score, box=None))
    return lines


def _extract_pdf_text(path: Path) -> str:
    import pypdfium2 as pdfium

    parts: list[str] = []
    doc = pdfium.PdfDocument(str(path))
    page_count = min(len(doc), MAX_PDF_PAGES)
    for i in range(page_count):
        page = doc[i]
        text_page = page.get_textpage()
        parts.append(text_page.get_text_bounded() or "")
    if len(doc) > MAX_PDF_PAGES:
        logger.warning("PDF 页数超过 %s，仅处理前 %s 页: %s", MAX_PDF_PAGES, MAX_PDF_PAGES, path)
    return "\n".join(parts)


def _ocr_pdf_pages(path: Path, lang: str, use_angle_cls: bool) -> list[OCRLine]:
    import pypdfium2 as pdfium

    all_lines: list[OCRLine] = []
    doc = pdfium.PdfDocument(str(path))
    page_count = min(len(doc), MAX_PDF_PAGES)
    for i in range(page_count):
        page = doc[i]
        bitmap = page.render(scale=PDF_RENDER_SCALE)
        pil_image = bitmap.to_pil()
        temp_path: Path | None = None
        try:
            with tempfile.NamedTemporaryFile(
                suffix=".png",
                prefix=f"ocr-pdf-p{i}-",
                delete=False,
            ) as tmp:
                pil_image.save(tmp.name, format="PNG")
                temp_path = Path(tmp.name)
            page_lines = run_ocr(
                OCRSettings(
                    image_path=temp_path,
                    lang=lang,
                    use_angle_cls=use_angle_cls,
                )
            )
            all_lines.extend(page_lines)
        finally:
            if temp_path is not None:
                try:
                    temp_path.unlink(missing_ok=True)
                except OSError:
                    pass
    return all_lines


def _extract_docx_text(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise ImportError(
            "python-docx 未安装。请执行: pip install python-docx"
        ) from exc

    doc = Document(str(path))
    parts: list[str] = []
    for para in doc.paragraphs:
        t = (para.text or "").strip()
        if t:
            parts.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [(c.text or "").strip() for c in row.cells]
            row_text = " | ".join(c for c in cells if c)
            if row_text:
                parts.append(row_text)
    return "\n".join(parts)


def recognize_resume_file(
    path: Path,
    lang: str,
    use_angle_cls: bool,
) -> tuple[list[OCRLine], dict[str, Any]]:
    """
    按文件类型识别简历内容。
    返回 (行列表, 元信息)，元信息含 mode: text_extract | ocr, file_type: image|pdf|docx。
    """
    resolved = Path(path).expanduser().resolve()
    if not resolved.exists():
        raise FileNotFoundError(f"File not found: {resolved}")

    suffix = resolved.suffix.lower()
    lang_norm = (lang or "ch").strip() or "ch"
    use_cls = bool(use_angle_cls)

    if suffix in ALLOWED_IMAGE_SUFFIX or not suffix:
        lines = run_ocr(
            OCRSettings(
                image_path=resolved,
                lang=lang_norm,
                use_angle_cls=use_cls,
            )
        )
        return lines, {"mode": "ocr", "file_type": "image"}

    if suffix == ".pdf":
        extracted = _extract_pdf_text(resolved)
        if _meaningful_char_count(extracted) >= MIN_PDF_TEXT_CHARS:
            lines = _lines_from_plain_text(extracted)
            logger.info(
                "PDF 文本抽取: path=%s chars=%s lines=%s",
                resolved,
                _meaningful_char_count(extracted),
                len(lines),
            )
            return lines, {"mode": "text_extract", "file_type": "pdf"}

        logger.info("PDF 文本过少，改分页 OCR: path=%s", resolved)
        lines = _ocr_pdf_pages(resolved, lang_norm, use_cls)
        return lines, {"mode": "ocr", "file_type": "pdf"}

    if suffix == ".docx":
        extracted = _extract_docx_text(resolved)
        lines = _lines_from_plain_text(extracted)
        if not lines:
            raise ValueError("未能从 Word 文档中提取到文字，请确认文件未加密且为有效 .docx")
        logger.info("DOCX 文本抽取: path=%s lines=%s", resolved, len(lines))
        return lines, {"mode": "text_extract", "file_type": "docx"}

    raise ValueError(
        f"不支持的文件类型: {suffix or '(无后缀)'}。"
        f"简历支持: 图片 {sorted(ALLOWED_IMAGE_SUFFIX)}、.pdf、.docx"
    )
